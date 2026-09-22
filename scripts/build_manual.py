"""Build the editable handbook using python-docx. Rendering is a separate QA step."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / '半导体研究生就业规划手册_我的模拟电路世界_v1.1.0.docx'
FONT = 'Microsoft YaHei'

def font_style(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');style.element.get_or_add_pPr().append(snap)

def cell_text(cell, text, header=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for para in cell.paragraphs:
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.line_spacing = Pt(13)
        for run in para.runs:
            run.font.size = Pt(9.4)
            run.font.bold = header
            run.font.color.rgb = RGBColor(255,255,255) if header else RGBColor(0,0,0)
            run.font.name = FONT
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)

def table(doc, block):
    cols = len(block['headers'])
    tab = doc.add_table(rows=1, cols=cols)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    total = 17.0
    ratios = {2:[0.26,0.74],3:[0.23,0.40,0.37]}.get(cols,[1/cols]*cols)
    for col, ratio in zip(tab.columns,ratios): col.width = Cm(total*ratio)
    for c, label in zip(tab.rows[0].cells,block['headers']): cell_text(c,label,True)
    rep = OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in block['rows']:
        cells=tab.add_row().cells
        for c, value in zip(cells,row): cell_text(c,value)
    pr=tab._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,value in [('top','75'),('bottom','75'),('left','90'),('right','90')]:
        el=OxmlElement('w:'+side);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    for i,row in enumerate(tab.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,c in enumerate(row.cells):
            c.width=Cm(total*ratios[j])
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243F5B' if i==0 else ('F1F5F8' if i%2==0 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shade)
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(1);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1
    spacer.add_run().font.size=Pt(2)

def hyperlink(para,label,url):
    rel=para.part.relate_to(url,RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'235A81');pr.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'19');pr.append(size)
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),FONT);fonts.set(qn('w:hAnsi'),FONT);fonts.set(qn('w:eastAsia'),FONT);pr.append(fonts)
    run.append(pr);txt=OxmlElement('w:t');txt.text=label;run.append(txt);link.append(run);para._p.append(link)

def main():
    source=json.loads((ROOT/'docs/manual-content.json').read_text(encoding='utf-8'))
    doc=Document()
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    sec.header_distance=Cm(0.7);sec.footer_distance=Cm(0.7);sec.different_first_page_header_footer=True
    normal=doc.styles['Normal'];font_style(normal,10.5)
    normal.paragraph_format.line_spacing=Pt(15)
    normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',28),('Subtitle',14),('Heading 1',18),('Heading 2',12)]:
        st=doc.styles[name];font_style(st,size,name!='Subtitle');st.paragraph_format.space_after=Pt(8);st.paragraph_format.space_before=Pt(9 if name=='Heading 2' else 0)
        st.paragraph_format.line_spacing=Pt(size*1.3)
        if name.startswith('Heading'):st.paragraph_format.keep_with_next=True
    for name in ['List Bullet','List Number']:
        st=doc.styles[name];font_style(st,10.5);st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=Pt(15)
    doc.core_properties.title=source['title'];doc.core_properties.author=source['author']
    doc.core_properties.subject='岗位能力驱动的半导体研究生学习与求职规划'
    doc.core_properties.keywords='半导体,研究生,就业规划,Skill,Codex,WorkBuddy'
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    for i,page in enumerate(source['pages']):
        if i:
            title=doc.add_paragraph(f'{i+1:02d}  {page["title"]}',style='Heading 1')
            title.paragraph_format.page_break_before=True
        else:
            title=doc.add_paragraph(page['title'],style='Title');title.paragraph_format.space_before=Pt(24)
            title.paragraph_format.space_after=Pt(16)
        for block in page['blocks']:
            typ=block['type']
            if typ=='table':table(doc,block)
            elif typ=='bullets':
                for item in block['items']:doc.add_paragraph(item,style='List Bullet')
            elif typ=='h2':doc.add_paragraph(block['text'],style='Heading 2')
            elif typ=='prompt':
                para=doc.add_paragraph(block['text'])
                para.paragraph_format.left_indent=Cm(0.35);para.paragraph_format.right_indent=Cm(0.2)
                for run in para.runs:run.font.size=Pt(10)
            elif typ=='links':
                para=doc.add_paragraph()
                para.paragraph_format.space_after=Pt(5)
                if block.get('prefix'):para.add_run(block['prefix'])
                for j,item in enumerate(block['items']):
                    if j:para.add_run('  |  ')
                    hyperlink(para,item['label'],item['url'])
            else:doc.add_paragraph(block['text'])
        if i==27:
            para=doc.add_paragraph();para.paragraph_format.space_after=Pt(2)
            links=[('OpenAI技能文档','https://learn.chatgpt.com/docs/build-skills'),('WorkBuddy技能规范','https://open.workbuddy.cn/docs/skill'),('MIT 6.012','https://ocw.mit.edu/courses/6-012-microelectronic-devices-and-circuits-fall-2009/'),('TI运放课程','https://www.ti.com/video/series/precision-labs/ti-precision-labs-op-amps.html')]
            for j,(label,url) in enumerate(links):
                if j:para.add_run('  |  ')
                hyperlink(para,label,url)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for para in doc.paragraphs:
        for border in list(para._p.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(json.dumps({'path':str(OUT),'planned_pages':len(source['pages']),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},ensure_ascii=False))

if __name__=='__main__':main()
